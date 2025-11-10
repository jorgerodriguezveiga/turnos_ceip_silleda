import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm
import calendar
from datetime import date
import argparse
import json
from bs4 import BeautifulSoup

DIAS_SEMANA = {0: "Luns", 1: "Martes", 2: "Mércores", 3: "Xoves", 4: "Venres"}
MESES = {
    1: "Xaneiro",
    2: "Febreiro",
    3: "Marzo",
    4: "Abril",
    5: "Maio",
    6: "Xunio",
    7: "Xullo",
    8: "Agosto",
    9: "Setembro",
    10: "Outubro",
    11: "Novembro",
    12: "Decembro",
}


def turnos_grupos(contador_zona, contador_baja, grupo_info) -> str:
    profesores = grupo_info["profesores"]
    num_profesores = len(profesores)

    # Rotaciones de zona
    fijos = grupo_info.get("fijos", {})
    profesores_fijos = list(fijos.keys())
    profesores_rotativos_zona = num_profesores - len(profesores_fijos)
    letras_fijas = list(fijos.values())
    letras = [i for i in grupo_info["zonas"] if i not in letras_fijas]
    indice = contador_zona % profesores_rotativos_zona

    # Rotaciones de baja
    no_baja = grupo_info.get("no_baja", [])
    posibles_profesores_baja = [p for p in profesores if p not in no_baja]
    profesor_baja = posibles_profesores_baja[contador_baja % len(posibles_profesores_baja)]

    orden_letras = letras[indice:] + letras[:indice]
    asignaciones = []
    contador_letra = 0
    for profesor in profesores:
        if profesor in fijos:
            texto = f"{profesor} ({fijos[profesor]})"
        else:
            texto = f"{profesor} ({orden_letras[contador_letra]})"
            contador_letra += 1

        if profesor == profesor_baja:
            texto = f"{texto} <b>baixa</b>"

        asignaciones.append(texto)

    return "\n".join(asignaciones)


def generar_fechas_laborales_mes(year, month) -> list[date]:
    # Genera una matriz de semanas (cada semana es una lista de días)
    month_days = calendar.monthcalendar(year, month)

    # Convierte a lista de fechas reales
    dates = [
        date(year, month, day)
        for week in month_days
        for day in week
        if day != 0 and date(year, month, day).weekday() not in [5, 6]
    ]
    return dates


def crear_asignaciones_mes(contador_grupo_inner, contador_baja, informacion, year, month):
    festivos = informacion["festivos"]
    grupos = informacion["grupos"]
    d_semana_grupos = {d: [g for g, info in grupos.items() if d in info["dias"]] for d in range(5)}
    fechas = generar_fechas_laborales_mes(year, month)
    asignaciones = []
    contador_grupo_outer = 0
    for fecha in fechas:
        grupo = None
        turnos = None
        if fecha.day not in festivos.get(f"{fecha.month}/{fecha.year}", []):
            d_semana = fecha.weekday()
            posibles_grupos = d_semana_grupos[d_semana]
            if len(posibles_grupos) == 1:
                grupo = posibles_grupos[0]
            else:
                grupo = posibles_grupos[contador_grupo_outer % len(posibles_grupos)]
                contador_grupo_outer += 1
            turnos = turnos_grupos(contador_grupo_inner[grupo], contador_baja[grupo], grupos[grupo])
            contador_grupo_inner[grupo] -= 1
            contador_baja[grupo] += 1

        asignaciones.append(
            {
                "fecha": fecha,
                "grupo": grupo,
                "turnos": turnos,
            }
        )

    return pd.DataFrame(asignaciones)


def misma_semana(f1, f2):
    return f1.isocalendar()[:2] == f2.isocalendar()[:2]


def set_cell_background(cell, hex_color):
    """
    Pone color de fondo a una celda.

    cell: celda de la tabla (docx.table._Cell)
    hex_color: color en formato hexadecimal, por ejemplo '#d9ead3'
    """
    # Convertir hex a RGB
    hex_color = hex_color.lstrip("#")  # quitar #
    r, g, b = tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))

    # Convertir a hexadecimal para el XML (sin #)
    hex_color_xml = "{:02X}{:02X}{:02X}".format(r, g, b)

    # Crear el shading XML
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), hex_color_xml))

    # Añadir shading a la celda
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(shading_elm)


def añadir_texto(p, texto, size=8):
    soup = BeautifulSoup(texto, "html.parser")
    for elem in soup:
        if elem.name == "b":
            run = p.add_run(elem.text)
            run.bold = True
        else:
            run = p.add_run(str(elem))

        run.font.size = Pt(size)


def crear_docx(calendarios, colores, motivo):
    # 2️⃣ Crear un nuevo documento Word
    doc = Document()
    section = doc.sections[0]

    section.top_margin = Cm(1)  # margen superior
    section.bottom_margin = Cm(1)  # margen inferior
    section.left_margin = Cm(2)  # margen izquierdo
    section.right_margin = Cm(2)  # margen derecho

    for mes, df in calendarios.items():
        heading = doc.add_heading(f"GARDAS {motivo.upper()} {MESES[mes].upper()}\n", level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 3️⃣ Añadir una tabla al documento con el tamaño adecuado
        table = doc.add_table(rows=1, cols=len(DIAS_SEMANA))
        table.style = "Table Grid"

        # 4️⃣ Agregar encabezados
        hdr_cells = table.rows[0].cells
        for i, dia in DIAS_SEMANA.items():
            hdr_cells[i].text = str(dia).upper()

        # 5️⃣ Agregar las filas del DataFrame
        fecha_anterior = None
        for _, row in df.iterrows():
            fecha: date = row["fecha"]
            if fecha_anterior is None or not misma_semana(fecha_anterior, fecha):
                row_cells = table.add_row().cells

            dia_semana = fecha.weekday()
            # Crear párrafo en la celda
            cell = row_cells[dia_semana]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Número del dia
            añadir_texto(p, f"<b>{fecha.day}</b>", 14)

            # Turnos
            turnos = row["turnos"]
            if turnos:
                añadir_texto(p, f"\n{turnos}", 8)

            grupo = row["grupo"]
            color = colores.get(grupo)
            if color:
                set_cell_background(cell, color)

            fecha_anterior = fecha

        # Nueva página
        doc.add_page_break()

    return doc


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generar documento Word")
    parser.add_argument("json_path", help="Ruta al archivo JSON con la información")
    parser.add_argument("--motivo", help="Nombre del motivo de la guardia")
    parser.add_argument("--output", default="GARDAS.docx", help="Nombre del archivo Word de salida")
    args = parser.parse_args()

    motivo = args.motivo

    # Cargar JSON
    with open(args.json_path, "r", encoding="utf-8") as f:
        informacion = json.load(f)

    grupos = informacion["grupos"]
    contador_grupo_inner = {grupo: info["contador"] for grupo, info in grupos.items()}
    contador_baja = {grupo: info["contador_baja"] for grupo, info in grupos.items()}
    meses_años = [[int(i) for i in my.split("/")] for my in informacion["festivos"].keys()]
    calendarios = {}
    for mes, año in meses_años:
        df_asignaciones = crear_asignaciones_mes(contador_grupo_inner, contador_baja, informacion, año, mes)
        calendarios[mes] = df_asignaciones

    colores = {grupo: info["color"] for grupo, info in grupos.items()}
    doc = crear_docx(calendarios, colores, motivo)

    # 6️⃣ Guardar el documento
    file = args.output
    doc.save(file)
    print(f"✅ Documento guardado como '{file}'")
